
from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt", ".md"}:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        return [Document(page_content=text, metadata={"source": str(path)})]
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return [Document(page_content=text, metadata={"source": str(path)})]

    raise ValueError(f"Unsupported file type: {suffix}. Supported types are: {SUPPORTED}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    docs = [d for d in docs if d.page_content.strip()]
    return splitter.split_documents(docs)